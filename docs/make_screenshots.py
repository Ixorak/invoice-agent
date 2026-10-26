"""Build the README screenshots from the output of a real demo run.

    python -m invoice_agent demo --width 150 --export-console docs/demo-console.svg --eval-report docs/eval.md
    python docs/make_screenshots.py

* docs/demo-console.svg is rich's recording of what the demo printed in the
  console. The two console pictures are crops of it.
* The Excel picture is drawn from demo/output/invoices.xlsx as it is on disk:
  values, number formats, fills, fonts, column widths and row heights are read
  with openpyxl and laid out as an HTML grid that looks like Excel.

Both are opened in a browser with Playwright (Microsoft Edge, else Chromium)
and saved as PNG files in docs/. Nothing in the pictures is typed by hand.
"""

from __future__ import annotations

import html
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from playwright.sync_api import sync_playwright

DOCS = Path(__file__).resolve().parent
ROOT = DOCS.parent
CONSOLE_SVG = DOCS / "demo-console.svg"
WORKBOOK = ROOT / "demo" / "output" / "invoices.xlsx"


def launch(p):
    try:
        return p.chromium.launch(channel="msedge")
    except Exception:  # Edge not installed: use Playwright's Chromium (playwright install chromium)
        return p.chromium.launch()


# --------------------------------------------------------------------------- console
CROP_JS = """
([startText, endText]) => {
  // From the line that contains startText to the first line below it that contains endText.
  const texts = [...document.querySelectorAll('svg text')];
  const norm = t => t.textContent.replace(/\\u00a0/g, ' ');
  const start = texts.find(t => norm(t).includes(startText));
  if (!start) return null;
  const s = start.getBoundingClientRect();
  const end = texts.find(t => norm(t).includes(endText) && t.getBoundingClientRect().y > s.y);
  if (!end) return null;
  const e = end.getBoundingClientRect();
  // Distance between two console lines, from the positions of all text elements.
  const tops = [...new Set(texts.map(t => Math.round(t.getBoundingClientRect().y)))].sort((a, b) => a - b);
  const gaps = tops.slice(1).map((y, i) => y - tops[i]).filter(d => d > 5).sort((a, b) => a - b);
  const pitch = gaps[Math.floor(gaps.length / 2)];
  const frame = document.querySelector('svg').getBoundingClientRect();
  const top = s.y - pitch * 0.5;
  return {x: frame.x, y: top, width: frame.width, height: e.y + pitch * 0.92 - top};
}
"""


def console_shots(browser) -> None:
    svg = CONSOLE_SVG.read_text(encoding="utf-8")
    page = browser.new_page(viewport={"width": 1900, "height": 1000}, device_scale_factor=1)
    page.set_content(f"<html><body style='margin:0;background:#0d1117'>{svg}</body></html>")
    page.wait_for_load_state("networkidle")
    page.evaluate("document.fonts.ready")
    width = page.evaluate("document.querySelector('svg').getBoundingClientRect().width")
    height = page.evaluate("document.querySelector('svg').getBoundingClientRect().height")
    page.set_viewport_size({"width": int(width) + 1, "height": int(height) + 1})

    crops = {
        "console-summary.png": ("invoice-agent: summary", "per invoice"),
        "console-evaluation.png": ("Evaluation: status and codes", "Time per invoice"),
    }
    for name, (start, end) in crops.items():
        clip = page.evaluate(CROP_JS, [start, end])
        if clip is None:
            sys.exit(f"Could not find '{start}' ... '{end}' in {CONSOLE_SVG.name}")
        page.screenshot(path=str(DOCS / name), clip=clip)
        print("wrote", DOCS / name)
    page.close()


# --------------------------------------------------------------------------- excel
def _display(value, number_format: str) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S" if "hh" in number_format else "%Y-%m-%d")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if number_format == "#,##0.00":
            return f"{value:,.2f}"
        if number_format == "0.0":
            return f"{value:.1f}"
        return f"{value:g}"
    return str(value)


def _rgb(color) -> str | None:
    rgb = getattr(color, "rgb", None)
    if isinstance(rgb, str) and len(rgb) == 8 and rgb != "00000000":
        return "#" + rgb[2:]
    return None


def sheet_html(path: Path, sheet: str, last_column: int) -> str:
    wb = load_workbook(path)
    ws = wb[sheet]
    widths = [round((ws.column_dimensions[get_column_letter(c)].width or 8.43) * 7 + 5)
              for c in range(1, last_column + 1)]
    freeze_col = 1  # freeze_panes is "B2": column A and row 1 are frozen
    cols = "".join(f"<col style='width:{w}px'>" for w in [42] + widths)
    head = "<tr class='letters'><th class='corner'></th>" + "".join(
        f"<th>{get_column_letter(c)}</th>" for c in range(1, last_column + 1)) + "</tr>"
    rows = []
    for r in range(1, ws.max_row + 1):
        height = ws.row_dimensions[r].height
        style = f" style='height:{round(height * 96 / 72)}px'" if height else ""
        cells = [f"<th class='rownum'>{r}</th>"]
        for c in range(1, last_column + 1):
            cell = ws.cell(row=r, column=c)
            css = []
            fill = _rgb(cell.fill.fgColor) if cell.fill and cell.fill.fill_type == "solid" else None
            if fill:
                css.append(f"background:{fill}")
            if cell.font is not None:
                if cell.font.b:
                    css.append("font-weight:600")
                color = _rgb(cell.font.color) if cell.font.color is not None else None
                if color:
                    css.append(f"color:{color}")
            numeric = isinstance(cell.value, (int, float, datetime, date)) and not isinstance(cell.value, bool)
            horizontal = cell.alignment.horizontal if cell.alignment else None
            css.append("text-align:" + (horizontal or ("right" if numeric else "left")))
            if cell.alignment is not None and cell.alignment.wrap_text:
                css.append("white-space:normal")
            classes = []
            if r == 1:
                classes.append("hdr")
            if c == freeze_col:
                classes.append("frozen-col")
            if r == 1:
                classes.append("frozen-row")
            text = html.escape(_display(cell.value, cell.number_format))
            button = "<span class='filter'>&#9662;</span>" if r == 1 else ""
            cls = f" class='{' '.join(classes)}'" if classes else ""
            cells.append(f"<td{cls} style='{';'.join(css)}'>{text}{button}</td>")
        rows.append(f"<tr{style}>" + "".join(cells) + "</tr>")
    tabs = "".join(
        f"<span class='tab{' active' if name == sheet else ''}'>{html.escape(name)}</span>" for name in wb.sheetnames)
    return f"""<!doctype html><html><head><meta charset='utf-8'><style>
      body {{ margin:0; background:#fff; font: 13px Calibri, 'Segoe UI', Arial, sans-serif; color:#000 }}
      table {{ border-collapse: collapse; table-layout: fixed; }}
      td, th {{ border:1px solid #e1e1e1; padding:2px 5px; height:20px; overflow:hidden; white-space:nowrap;
               vertical-align: top; text-overflow: clip; position: relative }}
      th {{ background:#f3f3f3; color:#444; font-weight:400; text-align:center; vertical-align: middle }}
      tr.letters th {{ height:20px }}
      th.rownum {{ text-align:center }}
      td.hdr {{ vertical-align: middle; padding-right: 20px; white-space: normal }}
      td.frozen-col {{ border-right: 2px solid #9e9e9e }}
      td.frozen-row {{ border-bottom: 2px solid #9e9e9e }}
      .filter {{ position:absolute; right:3px; top:50%; transform: translateY(-50%); width:13px; height:13px;
                 border:1px solid #b9b9b9; background:#fff; color:#444; font-size:9px; line-height:13px;
                 text-align:center; font-weight:400 }}
      .tabs {{ border-top:1px solid #d0d0d0; background:#f3f3f3; padding:0 0 0 48px; height:30px;
               display:flex; align-items:stretch }}
      .tab {{ padding:6px 16px; color:#444; border-right:1px solid #d6d6d6 }}
      .tab.active {{ background:#fff; color:#185c37; font-weight:600; border-bottom:3px solid #1e7145 }}
    </style></head><body>
      <table id='grid'><colgroup>{cols}</colgroup>{head}{''.join(rows)}</table>
      <div class='tabs'>{tabs}</div>
    </body></html>"""


def excel_shot(browser, sheet: str, last_column: int, name: str) -> None:
    page = browser.new_page(viewport={"width": 1700, "height": 600}, device_scale_factor=1)
    page.set_content(sheet_html(WORKBOOK, sheet, last_column=last_column))
    width = page.evaluate("document.getElementById('grid').getBoundingClientRect().width")
    height = page.evaluate("document.body.scrollHeight")
    page.set_viewport_size({"width": int(width) + 1, "height": int(height)})
    page.screenshot(path=str(DOCS / name), full_page=True)
    print("wrote", DOCS / name)
    page.close()


def main() -> None:
    for required in (CONSOLE_SVG, WORKBOOK):
        if not required.exists():
            sys.exit(f"Missing {required}. Run the demo first (see the docstring of this script).")
    with sync_playwright() as p:
        browser = launch(p)
        console_shots(browser)
        excel_shot(browser, "Invoices", last_column=12, name="excel-invoices.png")
        excel_shot(browser, "Review", last_column=7, name="excel-review.png")
        browser.close()


if __name__ == "__main__":
    main()
