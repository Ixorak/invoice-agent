"""output/invoices.xlsx: three sheets, appended to on every run.

* Invoices   - one row per invoice, with status and codes
* Line items - every line of every invoice
* Review     - one row per problem, only for invoices that need a person
"""

from __future__ import annotations

import math
import os
import tempfile
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from typing import TYPE_CHECKING

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .checks import CODE_TITLES, EXTRACTION_FAILED, OK
from .iban import format_iban

if TYPE_CHECKING:  # pragma: no cover
    from .pipeline import InvoiceResult


class ExcelLockedError(Exception):
    """The workbook is open in another program (usually Excel)."""


MONEY = "#,##0.00"
DATE = "yyyy-mm-dd"
DATETIME = "yyyy-mm-dd hh:mm:ss"
RATE = "0.0##"
QTY = "#,##0.###"
SECONDS = "0.0"

# (header, width, number format)
INVOICE_COLUMNS = [
    ("File", 34, None),
    ("Status", 16, None),
    ("Codes", 30, None),
    ("Supplier", 28, None),
    ("Invoice number", 18, None),
    ("Invoice date", 12, DATE),
    ("Due date", 12, DATE),
    ("Currency", 9, None),
    ("Subtotal", 12, MONEY),
    ("VAT", 11, MONEY),
    ("Total (printed)", 14, MONEY),
    ("Total (computed)", 15, MONEY),
    ("IBAN", 30, None),
    ("Supplier VAT ID", 22, None),
    ("Supplier address", 46, None),
    ("Payment terms", 36, None),
    ("Language", 9, None),
    ("Read from", 10, None),
    ("Line items", 10, None),
    ("Duration (s)", 11, SECONDS),
    ("Processed at", 19, DATETIME),
    ("Moved to", 44, None),
    ("JSON report", 48, None),
    ("SHA-256", 66, None),
]
LINE_COLUMNS = [
    ("File", 34, None),
    ("Supplier", 28, None),
    ("Invoice number", 18, None),
    ("Line", 6, None),
    ("Description", 44, None),
    ("Quantity", 10, QTY),
    ("Unit price", 12, MONEY),
    ("VAT rate (%)", 12, RATE),
    ("Amount", 12, MONEY),
    ("Currency", 9, None),
]
REVIEW_COLUMNS = [
    ("File", 34, None),
    ("Supplier", 28, None),
    ("Invoice number", 18, None),
    ("Status", 18, None),
    ("Code", 24, None),
    ("Problem", 50, None),
    ("Explanation", 90, None),
    ("Processed at", 19, DATETIME),
]
SHEETS = {"Invoices": INVOICE_COLUMNS, "Line items": LINE_COLUMNS, "Review": REVIEW_COLUMNS}

_HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_OK_FILL = PatternFill("solid", fgColor="C6EFCE")
_OK_FONT = Font(bold=True, color="006100")
_REVIEW_FILL = PatternFill("solid", fgColor="FFD8A8")
_REVIEW_FONT = Font(bold=True, color="8A4100")
_THIN = Side(style="thin", color="D0D7DE")


def _excel_value(value):
    if isinstance(value, Decimal):
        return float(value)
    return value


def _set(cell, value, number_format: str | None = None) -> None:
    value = _excel_value(value)
    if isinstance(value, str):
        value = ILLEGAL_CHARACTERS_RE.sub("", value)  # control characters are not allowed in xlsx
    cell.value = value
    if isinstance(value, str) and cell.data_type == "f":
        # Text from an invoice must never become a formula (e.g. "=HYPERLINK(...)").
        cell.data_type = "s"
    if number_format and value is not None:
        cell.number_format = number_format


def _style_status(cell) -> None:
    ok = cell.value == OK
    cell.fill = _OK_FILL if ok else _REVIEW_FILL
    cell.font = _OK_FONT if ok else _REVIEW_FONT
    cell.alignment = Alignment(horizontal="center", vertical="top")


def _new_workbook() -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    for title, columns in SHEETS.items():
        ws = wb.create_sheet(title)
        _write_header(ws, columns)
    return wb


def _write_header(ws, columns) -> None:
    for index, (header, width, _) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=index, value=header)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "B2"  # header row and file column stay visible
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}1"


def _append(ws, columns, values: list) -> int:
    row = ws.max_row + 1
    lines_needed = 1
    for index, ((_, width, fmt), value) in enumerate(zip(columns, values), start=1):
        cell = ws.cell(row=row, column=index)
        _set(cell, value, fmt)
        wrap = isinstance(value, str) and (len(value) > width or "\n" in value)
        cell.alignment = Alignment(vertical="top", wrap_text=wrap)
        cell.border = Border(bottom=_THIN)
        if wrap:
            per_line = max(int(width * 1.1) - 2, 1)
            lines = sum(max(1, math.ceil(len(part) / per_line)) for part in value.split("\n"))
            lines_needed = max(lines_needed, min(lines, 12))
    if lines_needed > 1:
        ws.row_dimensions[row].height = 15 * lines_needed  # Excel does not resize rows written by openpyxl
    return row


def _refresh_filter(ws, columns) -> None:
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max(ws.max_row, 1)}"


def _sheet(wb, title):
    if title in wb.sheetnames:
        return wb[title]
    ws = wb.create_sheet(title)
    _write_header(ws, SHEETS[title])
    return ws


def ensure_writable(path: Path) -> None:
    """Fail early (before calling the AI) if Excel has the workbook open."""
    if not path.exists():
        return
    try:
        with open(path, "r+b"):
            pass
    except PermissionError as exc:
        raise ExcelLockedError(f"{path.name} is open in another program. Close it and run again.") from exc


def _naive(dt: datetime) -> datetime:
    return dt.replace(tzinfo=None)


def invoice_row(r: "InvoiceResult") -> list:
    ex = r.extraction
    supplier = ex.supplier if ex else None
    info = ex.invoice if ex else None
    totals = ex.totals if ex else None
    return [
        r.file,
        r.status,
        ", ".join(r.codes),
        supplier.name if supplier else None,
        info.number if info else None,
        info.date if info else None,
        info.due_date if info else None,
        info.currency if info else None,
        totals.subtotal if totals else None,
        r.checks.vat_used if r.checks else None,
        totals.total_printed if totals else None,
        r.checks.computed_total if r.checks else None,
        format_iban(ex.payment.iban) if ex else None,
        supplier.vat_id if supplier else None,
        supplier.address if supplier else None,
        ex.payment.terms if ex else None,
        ex.language if ex else None,
        r.read_mode,
        len(ex.line_items) if ex else None,
        round(r.duration_seconds, 1),
        _naive(r.processed_at),
        r.moved_to_display or "not moved (still in the inbox)",
        r.json_display,
        r.sha256,
    ]


def _recorded_shas(ws) -> set[str]:
    """SHA-256 of the files that already have a final row (anything but EXTRACTION_FAILED)."""
    headers = [h for h, _, _ in INVOICE_COLUMNS]
    sha_index, status_index = headers.index("SHA-256"), headers.index("Status")
    found = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        if len(row) > sha_index and row[sha_index] and row[status_index] != EXTRACTION_FAILED:
            found.add(row[sha_index])
    return found


def append_results(path: Path, results: list["InvoiceResult"]) -> None:
    """Add the results at the end of the workbook (it is created if needed)."""
    if not results:
        return
    ensure_writable(path)
    wb = load_workbook(path) if path.exists() else _new_workbook()
    invoices, lines, review = _sheet(wb, "Invoices"), _sheet(wb, "Line items"), _sheet(wb, "Review")
    already = _recorded_shas(invoices)

    for r in results:
        if r.sha256 in already:
            continue  # e.g. a run interrupted right after this step: a file is never added twice
        row = _append(invoices, INVOICE_COLUMNS, invoice_row(r))
        _style_status(invoices.cell(row=row, column=2))

        ex = r.extraction
        if ex:
            for number, line in enumerate(ex.line_items, start=1):
                _append(lines, LINE_COLUMNS, [
                    r.file, ex.supplier.name, ex.invoice.number, number, line.description,
                    line.quantity, line.unit_price, line.vat_rate, line.amount, ex.invoice.currency,
                ])

        if r.status != OK:
            for issue in r.issues:
                row = _append(review, REVIEW_COLUMNS, [
                    r.file,
                    ex.supplier.name if ex else None,
                    ex.invoice.number if ex else None,
                    r.status,
                    issue.code,
                    CODE_TITLES.get(issue.code, issue.code),
                    issue.message,
                    _naive(r.processed_at),
                ])
                _style_status(review.cell(row=row, column=4))

    for ws, columns in ((invoices, INVOICE_COLUMNS), (lines, LINE_COLUMNS), (review, REVIEW_COLUMNS)):
        _refresh_filter(ws, columns)

    path.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp = tempfile.mkstemp(prefix=".invoices-", suffix=".xlsx", dir=path.parent)
    os.close(fd)
    try:
        wb.save(tmp)
        os.replace(tmp, path)
    except PermissionError as exc:
        Path(tmp).unlink(missing_ok=True)
        raise ExcelLockedError(f"{path.name} is open in another program. Close it and run again.") from exc
    except BaseException:
        Path(tmp).unlink(missing_ok=True)
        raise
