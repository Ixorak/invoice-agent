from openpyxl import load_workbook

from invoice_agent.excel import SHEETS
from invoice_agent.pipeline import Agent

from .conftest import FakeBackend, copy_sample, payload


def test_workbook_layout_and_append(config, workspace):
    copy_sample("03-nordwerk-demo.pdf", workspace.inbox, "a.pdf")
    Agent(config, workspace, FakeBackend(payload())).process_inbox()
    copy_sample("04-tipografia-esempio.pdf", workspace.inbox, "b.pdf")
    Agent(config, workspace, FakeBackend(payload({"invoice.number": "R-1002", "totals.total_printed": "1"}))).process_inbox()

    wb = load_workbook(workspace.excel_path)
    assert wb.sheetnames == list(SHEETS)
    invoices = wb["Invoices"]
    assert invoices.max_row == 3                       # header + 2 runs appended
    assert invoices.freeze_panes == "B2"
    assert invoices.auto_filter.ref == "A1:X3"
    assert [c.value for c in invoices[1]][:3] == ["File", "Status", "Codes"]
    assert invoices["B2"].value == "OK" and invoices["B2"].fill.fgColor.rgb.endswith("C6EFCE")   # green
    assert invoices["B3"].value == "NEEDS_REVIEW" and invoices["B3"].fill.fgColor.rgb.endswith("FFD8A8")  # orange
    assert invoices["I2"].number_format == "#,##0.00"   # subtotal
    assert invoices["F2"].number_format == "yyyy-mm-dd"  # invoice date
    assert wb["Review"].max_row == 2
    assert wb["Line items"].max_row == 5


def test_text_from_an_invoice_never_becomes_a_formula(config, workspace):
    copy_sample("03-nordwerk-demo.pdf", workspace.inbox, "a.pdf")
    evil = payload({"supplier.name": "=HYPERLINK(\"http://example.invalid\",\"click\")"})
    Agent(config, workspace, FakeBackend(evil)).process_inbox()
    cell = load_workbook(workspace.excel_path)["Invoices"]["D2"]
    assert cell.data_type == "s"
    assert cell.value.startswith("=HYPERLINK")


def test_control_characters_are_removed_before_writing():
    from openpyxl import Workbook

    from invoice_agent.excel import _set

    cell = Workbook().active["A1"]
    _set(cell, "excerpt\x01 with\x02 control characters")
    assert cell.value == "excerpt with control characters"


def test_the_same_file_is_never_added_twice(config, workspace):
    from invoice_agent.excel import append_results

    copy_sample("03-nordwerk-demo.pdf", workspace.inbox, "a.pdf")
    [result] = Agent(config, workspace, FakeBackend(payload())).process_inbox().results
    append_results(workspace.excel_path, [result])  # e.g. a run interrupted before the registry was saved
    wb = load_workbook(workspace.excel_path)
    assert wb["Invoices"].max_row == 2 and wb["Line items"].max_row == 3
