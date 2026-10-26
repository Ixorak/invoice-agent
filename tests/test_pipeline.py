"""The whole pipeline with a fake AI: folders, Excel, JSON, registry, skipping, duplicates."""

import json
from pathlib import Path

from openpyxl import load_workbook

from invoice_agent.backends import BackendError
from invoice_agent.pipeline import SKIPPED, Agent, list_pdfs

from .conftest import FakeBackend, copy_sample, payload


def _agent(config, workspace, *answers):
    return Agent(config, workspace, FakeBackend(*answers))


def _rows(path: Path, sheet: str) -> list[tuple]:
    wb = load_workbook(path)
    return list(wb[sheet].iter_rows(min_row=2, values_only=True))


def test_ok_invoice_goes_to_processed_with_json_and_excel(config, workspace):
    copy_sample("03-nordwerk-demo.pdf", workspace.inbox, "a.pdf")
    batch = _agent(config, workspace, payload()).process_inbox()
    assert batch.error is None
    [result] = batch.results
    assert result.status == "OK"
    assert (workspace.processed / "a.pdf").exists() and not (workspace.inbox / "a.pdf").exists()
    report = json.loads((workspace.json_dir / "a.json").read_text(encoding="utf-8"))
    assert report["status"] == "OK" and report["extraction"]["invoice"]["number"] == "R-1001"
    assert report["duration_seconds"] >= 0 and report["checks"]["computed_total"] == 270.25
    assert len(_rows(workspace.excel_path, "Invoices")) == 1
    assert len(_rows(workspace.excel_path, "Line items")) == 2
    assert _rows(workspace.excel_path, "Review") == []


def test_invoice_with_a_problem_goes_to_review_with_explanations(config, workspace):
    copy_sample("03-nordwerk-demo.pdf", workspace.inbox, "b.pdf")
    batch = _agent(config, workspace, payload({"totals.total_printed": "280.25"})).process_inbox()
    assert batch.results[0].status == "NEEDS_REVIEW"
    assert (workspace.review / "b.pdf").exists()
    review = _rows(workspace.excel_path, "Review")
    assert len(review) == 1
    assert review[0][4] == "TOTAL_MISMATCH" and "printed total" in review[0][6]


def test_same_file_twice_is_skipped_not_added_twice(config, workspace):
    copy_sample("03-nordwerk-demo.pdf", workspace.inbox, "c.pdf")
    _agent(config, workspace, payload()).process_inbox()
    # the very same file arrives again, under another name, in a later run
    copy_sample("03-nordwerk-demo.pdf", workspace.inbox, "c-again.pdf")
    backend = FakeBackend(payload())
    batch = Agent(config, workspace, backend).process_inbox()
    assert [r.status for r in batch.results] == [SKIPPED]
    assert backend.requests == []                      # no AI call
    assert len(_rows(workspace.excel_path, "Invoices")) == 1
    assert (workspace.inbox / "c-again.pdf").exists()  # ignored: left where it is


def test_other_file_with_same_supplier_and_number_is_a_duplicate(config, workspace):
    copy_sample("03-nordwerk-demo.pdf", workspace.inbox, "first.pdf")
    _agent(config, workspace, payload()).process_inbox()
    copy_sample("04-tipografia-esempio.pdf", workspace.inbox, "second.pdf")  # different bytes
    batch = _agent(config, workspace, payload()).process_inbox()           # a new run
    [result] = batch.results
    assert result.codes == ["DUPLICATE"]
    assert "first.pdf" in result.issues[0].message
    assert len(_rows(workspace.excel_path, "Invoices")) == 2               # appended


def test_backend_failure_stops_the_run_and_leaves_the_file(config, workspace):
    copy_sample("03-nordwerk-demo.pdf", workspace.inbox, "d.pdf")
    copy_sample("04-tipografia-esempio.pdf", workspace.inbox, "e.pdf")
    batch = _agent(config, workspace, BackendError("Claude Code is not logged in")).process_inbox()
    assert batch.error and "not logged in" in batch.error
    assert batch.results == []
    assert {p.name for p in list_pdfs(workspace.inbox)} == {"d.pdf", "e.pdf"}
    assert not workspace.excel_path.exists()
    assert not workspace.registry_path.exists()


def test_extraction_failed_goes_to_review(config, workspace):
    copy_sample("03-nordwerk-demo.pdf", workspace.inbox, "f.pdf")
    batch = _agent(config, workspace, "not json").process_inbox()
    [result] = batch.results
    assert result.status == "EXTRACTION_FAILED"
    assert (workspace.review / "f.pdf").exists()
    assert _rows(workspace.excel_path, "Invoices")[0][1] == "EXTRACTION_FAILED"
    assert _rows(workspace.excel_path, "Review")[0][4] == "EXTRACTION_FAILED"


def test_a_file_that_is_not_a_pdf_fails_cleanly(config, workspace):
    (workspace.inbox / "broken.pdf").write_bytes(b"this is not a pdf")
    backend = FakeBackend(payload())
    batch = Agent(config, workspace, backend).process_inbox()
    assert batch.results[0].status == "EXTRACTION_FAILED"
    assert backend.requests == []
    assert (workspace.review / "broken.pdf").exists()


def test_scanned_sample_is_sent_as_a_file(config, workspace):
    copy_sample("09-menuiserie-specimen-scan.pdf", workspace.inbox, "scan.pdf")
    backend = FakeBackend(payload())
    batch = Agent(config, workspace, backend).process_inbox()
    assert batch.results[0].read_mode == "scan"
    assert backend.requests[0].pdf_path is not None


def test_pdf_bytes_are_never_modified(config, workspace):
    source = copy_sample("03-nordwerk-demo.pdf", workspace.inbox, "g.pdf")
    before = source.read_bytes()
    _agent(config, workspace, payload({"totals.total_printed": "1.00"})).process_inbox()
    assert (workspace.review / "g.pdf").read_bytes() == before


def test_name_collision_in_destination_keeps_both_files(config, workspace):
    (workspace.processed / "h.pdf").write_bytes(b"older file with the same name")
    copy_sample("03-nordwerk-demo.pdf", workspace.inbox, "h.pdf")
    batch = _agent(config, workspace, payload()).process_inbox()
    moved = batch.results[0].moved_to
    assert moved.name.startswith("h-") and moved.exists()
    assert (workspace.processed / "h.pdf").read_bytes() == b"older file with the same name"


def test_injection_in_the_pdf_is_flagged_even_if_the_ai_misses_it(config, workspace):
    # Sample 10 contains a hidden instruction. Here the fake AI reports nothing:
    # the keyword scan of the PDF text must still flag the invoice.
    copy_sample("10-conseil-fantome.pdf", workspace.inbox, "i.pdf")
    batch = _agent(config, workspace, payload()).process_inbox()
    [result] = batch.results
    assert "SUSPICIOUS_INSTRUCTION" in result.codes
    assert result.status == "NEEDS_REVIEW"
